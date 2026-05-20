import logging
import os
import uuid
import io
import csv
from datetime import date, datetime
from fastapi import APIRouter, HTTPException, Query, status, Depends, UploadFile, File, Form
from fastapi.responses import FileResponse, StreamingResponse
from typing import List, Optional
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func as sqlfunc
from database import get_db
from sql_models import Gasto, Proveedor, Insumo, Categoria, Subcategoria, AbonoGasto, MedioPago, GastoAdjunto, Usuario, ParametroSistema
from models import Expense, ExpenseCreate, GastoAdjuntoResponse
from auth import get_current_user, require_internal_permission
from audit import auditar, snapshot
from fastapi import Request

router = APIRouter()
logger = logging.getLogger(__name__)

# Storage de adjuntos
# - Default: <backend>/uploads/expenses
# - Configurable via parametros_sistema.clave='gastos_adjuntos_directorio'
DEFAULT_UPLOAD_ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "uploads", "expenses")
os.makedirs(DEFAULT_UPLOAD_ROOT, exist_ok=True)
PARAM_DIR_KEY = "gastos_adjuntos_directorio"
MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB
ALLOWED_MIME_PREFIXES = ("image/", "application/pdf", "text/plain")


def _get_upload_dir(db: Session) -> str:
    """
    Devuelve el directorio absoluto donde se guardan adjuntos.
    Lee `parametros_sistema.gastos_adjuntos_directorio`; si está vacío, usa DEFAULT_UPLOAD_ROOT.
    Valida que sea absoluto, sin '..' y crea el directorio si no existe.
    """
    param = db.query(ParametroSistema).filter(ParametroSistema.clave == PARAM_DIR_KEY).first()
    configured = (param.valor or "").strip() if param else ""
    if not configured:
        return DEFAULT_UPLOAD_ROOT

    # Validaciones de seguridad
    if ".." in configured.split(os.sep):
        raise HTTPException(500,
            f"Parámetro '{PARAM_DIR_KEY}' inválido: no se permiten segmentos '..'. "
            f"Configurarlo en /admin/parametros → Configuración.")
    if not os.path.isabs(configured):
        raise HTTPException(500,
            f"Parámetro '{PARAM_DIR_KEY}' debe ser una ruta absoluta (comenzar con '/'). "
            f"Valor actual: '{configured}'. Configurarlo en /admin/parametros → Configuración.")

    # Asegurar que existe y es escribible
    try:
        os.makedirs(configured, exist_ok=True)
    except OSError as e:
        raise HTTPException(500,
            f"No se puede crear el directorio '{configured}': {e}. "
            f"Verifica permisos o cambia el valor en /admin/parametros → Configuración.")
    if not os.access(configured, os.W_OK):
        raise HTTPException(500,
            f"El directorio '{configured}' no es escribible por el backend.")
    return configured


def _resolve_adjunto_path(ruta_archivo: str) -> str:
    """
    Devuelve la ruta absoluta del archivo en disco.
      • Si ruta_archivo ya es absoluta (uploads nuevos con dir configurable): se usa tal cual.
      • Si es relativa (uploads legacy 'expenses/YYYY/MM/uuid.ext'): se resuelve bajo <backend>/uploads/.
    """
    if os.path.isabs(ruta_archivo):
        return ruta_archivo
    legacy_base = os.path.dirname(DEFAULT_UPLOAD_ROOT)   # <backend>/uploads
    return os.path.join(legacy_base, ruta_archivo)

SAFE_FIELDS = {"valor", "fecha", "categoria_id", "subcategoria_id",
               "proveedor_id", "medio_pago_id", "insumo_id",
               "cantidad_insumo", "estado", "tipo_costo"}


def _expense_dict(e: Gasto) -> dict:
    monto_pagado = float(e.monto_pagado or 0)
    valor        = float(e.valor or 0)
    return {
        "id": e.id,
        "valor": valor,
        "fecha": e.fecha,
        "categoria_id": e.categoria_id,
        "categoria_nombre": e.categoria_rel.nombre if e.categoria_rel else None,
        "subcategoria_id": e.subcategoria_id,
        "subcategoria_nombre": e.subcategoria_rel.nombre if e.subcategoria_rel else None,
        "proveedor_id": e.proveedor_id,
        "medio_pago_id": e.medio_pago_id,
        "insumo_id": e.insumo_id,
        "cantidad_insumo": float(e.cantidad_insumo) if e.cantidad_insumo else None,
        "insumo_nombre": e.insumo.nombre if e.insumo else None,
        "proveedor_nombre": e.proveedor.nombre if e.proveedor else None,
        "estado": e.estado or 'activo',
        "monto_pagado": monto_pagado,
        "saldo_pendiente": round(valor - monto_pagado, 2),
        "estado_pago": e.estado_pago or 'pendiente',
        # Auditoría
        "created_by": e.created_by,
        "created_by_nombre": e.creador.nombre if e.creador else None,
        "updated_at": e.updated_at,
        "anulado_at": e.anulado_at,
    }


def _validate_subcategoria_coherente(data: dict, db: Session):
    """La subcategoría debe pertenecer a la categoría seleccionada."""
    sub_id = data.get("subcategoria_id")
    cat_id = data.get("categoria_id")
    if not sub_id:
        return
    sub = db.query(Subcategoria).filter(Subcategoria.id == sub_id).first()
    if not sub:
        raise HTTPException(status_code=404, detail="Subcategoría no encontrada")
    if cat_id and sub.categoria_id != cat_id:
        raise HTTPException(
            status_code=400,
            detail=f"La subcategoría '{sub.nombre}' no pertenece a la categoría seleccionada",
        )
    # Si no vino categoria_id explícito, la inferimos de la subcategoría
    if not cat_id:
        data["categoria_id"] = sub.categoria_id


class AbonoCreate(BaseModel):
    monto:         float = Field(..., gt=0)
    fecha:         str   = Field(default_factory=lambda: str(date.today()))
    medio_pago_id: Optional[int] = None
    notas:         Optional[str] = None


def _load(expense_id: int, db: Session) -> Gasto:
    return (
        db.query(Gasto)
        .options(
            joinedload(Gasto.proveedor),
            joinedload(Gasto.insumo),
            joinedload(Gasto.categoria_rel),
            joinedload(Gasto.subcategoria_rel),
            joinedload(Gasto.creador),
        )
        .filter(Gasto.id == expense_id)
        .first()
    )


def _cat_dict(c: Categoria, *, include_subs: bool = True) -> dict:
    d = {
        "id": c.id,
        "nombre": c.nombre,
        "tipo": c.tipo,
        "tipo_costo": c.tipo_costo,
        "icono": c.icono,
        "color": c.color,
        "activo": c.activo,
        "orden": c.orden,
    }
    if include_subs:
        d["subcategorias"] = [
            {"id": s.id, "nombre": s.nombre, "icono": s.icono, "color": s.color,
             "activo": s.activo, "orden": s.orden}
            for s in (c.subcategorias or [])
            if s.activo
        ]
    return d


@router.get("/categorias", dependencies=[Depends(require_internal_permission("gastos.ver"))])
def get_categorias(
    only_active: bool = False,
    db: Session = Depends(get_db),
):
    """
    Listado consumido por la UI de Gastos. Filtra automáticamente al módulo 'gastos'
    e incluye subcategorías activas embebidas.
    """
    from sqlalchemy.orm import selectinload
    q = (
        db.query(Categoria)
        .options(selectinload(Categoria.subcategorias))
        .filter(Categoria.modulo_codigo == "gastos")
        .order_by(Categoria.orden, Categoria.id)
    )
    if only_active:
        q = q.filter(Categoria.activo == True)  # noqa: E712
    return [_cat_dict(c) for c in q.all()]


@router.patch("/categorias/{cat_id}/tipo-costo", dependencies=[Depends(require_internal_permission("gastos.editar"))])
def update_tipo_costo(cat_id: int, payload: dict, db: Session = Depends(get_db)):
    """Actualiza tipo_costo (directo|indirecto|null) de una categoría del módulo 'gastos'."""
    cat = db.query(Categoria).filter(
        Categoria.id == cat_id,
        Categoria.modulo_codigo == "gastos",
    ).first()
    if not cat:
        raise HTTPException(status_code=404, detail="Categoría no encontrada")
    nuevo = payload.get("tipo_costo")
    if nuevo not in ("directo", "indirecto", None):
        raise HTTPException(status_code=400, detail="tipo_costo debe ser 'directo', 'indirecto' o null")
    cat.tipo_costo = nuevo
    db.commit()
    return _cat_dict(cat)


@router.get("/cuentas-por-pagar", dependencies=[Depends(require_internal_permission("gastos.ver"))])
def get_cuentas_por_pagar(db: Session = Depends(get_db)):
    """Gastos activos con saldo pendiente, agrupados por proveedor."""
    gastos = (
        db.query(Gasto)
        .options(joinedload(Gasto.proveedor), joinedload(Gasto.categoria_rel))
        .filter(
            Gasto.estado != 'anulado',
            Gasto.proveedor_id.isnot(None),
            Gasto.estado_pago != 'pagado',
        )
        .order_by(Gasto.fecha.asc(), Gasto.id.asc())
        .all()
    )
    proveedores: dict = {}
    for g in gastos:
        pid = g.proveedor_id
        saldo = float(g.valor or 0) - float(g.monto_pagado or 0)
        if pid not in proveedores:
            proveedores[pid] = {
                'proveedor_id': pid,
                'nombre': g.proveedor.nombre if g.proveedor else f'Proveedor #{pid}',
                'total_deuda': 0.0,
                'facturas_pendientes': 0,
                'fecha_mas_antigua': str(g.fecha)[:10] if g.fecha else None,
                'facturas': [],
            }
        proveedores[pid]['total_deuda'] = round(proveedores[pid]['total_deuda'] + saldo, 2)
        proveedores[pid]['facturas_pendientes'] += 1
        fecha_g = str(g.fecha)[:10] if g.fecha else None
        if fecha_g and (not proveedores[pid]['fecha_mas_antigua'] or fecha_g < proveedores[pid]['fecha_mas_antigua']):
            proveedores[pid]['fecha_mas_antigua'] = fecha_g
        proveedores[pid]['facturas'].append(_expense_dict(g))
    result = sorted(proveedores.values(), key=lambda x: x['total_deuda'], reverse=True)
    return result


@router.get("/", response_model=List[Expense], dependencies=[Depends(require_internal_permission("gastos.ver"))])
def get_expenses(
    start_date:      Optional[str] = None,
    end_date:        Optional[str] = None,
    categoria_id:    Optional[int] = None,
    subcategoria_id: Optional[int] = None,
    proveedor_id:    Optional[int] = None,
    medio_pago_id:   Optional[int] = None,
    estado:          Optional[str] = None,       # activo | anulado
    estado_pago:     Optional[str] = None,       # pendiente | parcial | pagado
    search:          Optional[str] = None,       # busca en proveedor.nombre / categoria / subcategoria
    skip:  int = Query(0,   ge=0),
    limit: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db)
):
    query = (
        db.query(Gasto)
        .options(
            joinedload(Gasto.proveedor),
            joinedload(Gasto.insumo),
            joinedload(Gasto.categoria_rel),
            joinedload(Gasto.subcategoria_rel),
        )
        .order_by(Gasto.fecha.desc(), Gasto.id.desc())
    )
    if start_date and end_date:
        query = query.filter(Gasto.fecha >= start_date, Gasto.fecha <= end_date)
    if categoria_id is not None:
        query = query.filter(Gasto.categoria_id == categoria_id)
    if subcategoria_id is not None:
        query = query.filter(Gasto.subcategoria_id == subcategoria_id)
    if proveedor_id is not None:
        query = query.filter(Gasto.proveedor_id == proveedor_id)
    if medio_pago_id is not None:
        query = query.filter(Gasto.medio_pago_id == medio_pago_id)
    if estado:
        query = query.filter(Gasto.estado == estado)
    if estado_pago:
        query = query.filter(Gasto.estado_pago == estado_pago)
    if search:
        like = f"%{search.lower()}%"
        # Búsqueda case-insensitive en proveedor + categoría + subcategoría
        from sqlalchemy import or_, func as sqlfunc
        query = (
            query
            .outerjoin(Gasto.proveedor)
            .outerjoin(Gasto.categoria_rel)
            .outerjoin(Gasto.subcategoria_rel)
            .filter(or_(
                sqlfunc.lower(Proveedor.nombre).like(like),
                sqlfunc.lower(Categoria.nombre).like(like),
                sqlfunc.lower(Subcategoria.nombre).like(like),
            ))
        )
    return [_expense_dict(e) for e in query.offset(skip).limit(limit).all()]


@router.post("/", response_model=Expense, status_code=status.HTTP_201_CREATED, dependencies=[Depends(require_internal_permission("gastos.crear"))])
def create_expense(expense: ExpenseCreate, db: Session = Depends(get_db),
                   user: Usuario = Depends(get_current_user)):
    try:
        data = {k: v for k, v in expense.model_dump().items() if k in SAFE_FIELDS and v is not None}
        _validate_insumo(data, db)
        _validate_subcategoria_coherente(data, db)
        # Auto-propagar tipo_costo desde la categoría si no viene explícito
        if "tipo_costo" not in data and data.get("categoria_id"):
            cat = db.query(Categoria).filter(Categoria.id == data["categoria_id"]).first()
            if cat and cat.tipo_costo:
                data["tipo_costo"] = cat.tipo_costo
        # Estado de pago al crear: si el usuario marca 'pagado' en el formulario,
        # el gasto nace saldado (monto_pagado = valor). Cualquier otro valor (credito,
        # pendiente, None) deja los defaults de BD: monto_pagado=0, estado_pago='pendiente'.
        if expense.estado_pago == 'pagado':
            data["estado_pago"]  = 'pagado'
            data["monto_pagado"] = data.get("valor", 0)
        # Auditoría
        data["created_by"] = user.id
        db_expense = Gasto(**data)
        db.add(db_expense)
        db.commit()
        # Trigger fn_gasto_actualiza_insumo: actualiza stock, upserta
        # insumo_proveedores y recalcula costo_unitario como promedio ponderado.
        return _expense_dict(_load(db_expense.id, db))
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        logger.error("Error creating expense: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail="Error al registrar el gasto")


@router.put("/{expense_id}", response_model=Expense, dependencies=[Depends(require_internal_permission("gastos.editar"))])
def update_expense(expense_id: int, expense_update: ExpenseCreate, db: Session = Depends(get_db),
                   user: Usuario = Depends(get_current_user)):
    db_expense = db.query(Gasto).filter(Gasto.id == expense_id).first()
    if not db_expense:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")

    data = expense_update.model_dump(exclude_unset=True)
    _validate_insumo(data, db)
    # Para validar coherencia subcategoría↔categoría, considera el estado actual del registro
    merged = {
        "categoria_id":    data.get("categoria_id",    db_expense.categoria_id),
        "subcategoria_id": data.get("subcategoria_id", db_expense.subcategoria_id),
    }
    _validate_subcategoria_coherente(merged, db)
    if "categoria_id" not in data and merged.get("categoria_id") != db_expense.categoria_id:
        # _validate_subcategoria_coherente puede haber inferido categoria_id desde la sub
        data["categoria_id"] = merged["categoria_id"]

    for key, value in data.items():
        if key in SAFE_FIELDS:
            setattr(db_expense, key, value)

    # Auditoría
    db_expense.updated_by = user.id
    db_expense.updated_at = datetime.utcnow()

    db.commit()
    # Trigger fn_gasto_actualiza_insumo: ajusta inventario y recalcula
    # costo_unitario como promedio ponderado sobre todos los gastos activos.
    return _expense_dict(_load(db_expense.id, db))


@router.patch("/{expense_id}/anular")
def anular_expense(
    expense_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: Usuario = Depends(require_internal_permission("gastos.eliminar")),
):
    """Anula un gasto activo. Revierte el stock del insumo si aplica (vía trigger)
    y elimina los abonos asociados para liberar la caja del medio de pago."""
    db_expense = db.query(Gasto).filter(Gasto.id == expense_id).first()
    if not db_expense:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    if db_expense.estado == 'anulado':
        raise HTTPException(status_code=400, detail="El gasto ya está anulado")

    antes = snapshot(db_expense)
    # Eliminar abonos: el trigger trg_sync_estado_pago_gasto recalcula
    # monto_pagado/estado_pago del gasto. Sin esto, los abonos quedan
    # huérfanos y siguen restando del saldo del medio de pago.
    db.query(AbonoGasto).filter(AbonoGasto.gasto_id == expense_id).delete(
        synchronize_session=False
    )
    db_expense.estado = 'anulado'
    db_expense.anulado_by = user.id
    db_expense.anulado_at = datetime.utcnow()
    db.commit()

    auditar(db, user, "gasto", expense_id, "anular",
            antes=antes, despues=snapshot(db_expense), request=request)

    return _expense_dict(_load(db_expense.id, db))


@router.patch("/{expense_id}/reactivar", dependencies=[Depends(require_internal_permission("gastos.editar"))])
def reactivar_expense(expense_id: int, db: Session = Depends(get_db),
                      user: Usuario = Depends(get_current_user)):
    """Reactiva un gasto anulado. Re-suma el stock del insumo si aplica (vía trigger)."""
    db_expense = db.query(Gasto).filter(Gasto.id == expense_id).first()
    if not db_expense:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    if db_expense.estado != 'anulado':
        raise HTTPException(status_code=400, detail="El gasto no está anulado")
    db_expense.estado = 'activo'
    db_expense.anulado_by = None
    db_expense.anulado_at = None
    db_expense.updated_by = user.id
    db_expense.updated_at = datetime.utcnow()
    db.commit()
    return _expense_dict(_load(db_expense.id, db))


@router.get("/{expense_id}/abonos", dependencies=[Depends(require_internal_permission("gastos.ver"))])
def get_abonos(expense_id: int, db: Session = Depends(get_db)):
    gasto = db.query(Gasto).filter(Gasto.id == expense_id).first()
    if not gasto:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    abonos = (
        db.query(AbonoGasto)
        .options(joinedload(AbonoGasto.medio_pago))
        .filter(AbonoGasto.gasto_id == expense_id)
        .order_by(AbonoGasto.fecha.asc(), AbonoGasto.id.asc())
        .all()
    )
    return [
        {
            "id": a.id,
            "monto": float(a.monto),
            "fecha": a.fecha,
            "medio_pago_id": a.medio_pago_id,
            "medio_pago_nombre": a.medio_pago.nombre if a.medio_pago else None,
            "notas": a.notas,
        }
        for a in abonos
    ]


@router.post("/{expense_id}/abonos", status_code=201, dependencies=[Depends(require_internal_permission("gastos.editar"))])
def crear_abono(expense_id: int, body: AbonoCreate, db: Session = Depends(get_db)):
    gasto = db.query(Gasto).filter(Gasto.id == expense_id).first()
    if not gasto:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    if gasto.estado == 'anulado':
        raise HTTPException(status_code=400, detail="No se puede abonar a un gasto anulado")
    saldo = float(gasto.valor) - float(gasto.monto_pagado or 0)
    if body.monto > saldo + 0.01:
        raise HTTPException(status_code=400, detail=f"El abono ({body.monto}) supera el saldo pendiente ({round(saldo, 2)})")
    try:
        fecha_dt = date.fromisoformat(body.fecha)
    except ValueError:
        raise HTTPException(status_code=400, detail="Fecha inválida")
    abono = AbonoGasto(
        gasto_id=expense_id,
        monto=body.monto,
        fecha=fecha_dt,
        medio_pago_id=body.medio_pago_id,
        notas=body.notas,
    )
    db.add(abono)
    db.commit()
    db.refresh(gasto)
    return {"ok": True, "saldo_pendiente": round(float(gasto.valor) - float(gasto.monto_pagado), 2), "estado_pago": gasto.estado_pago}


@router.delete("/{expense_id}/abonos/{abono_id}", dependencies=[Depends(require_internal_permission("gastos.editar"))])
def eliminar_abono(expense_id: int, abono_id: int, db: Session = Depends(get_db)):
    abono = db.query(AbonoGasto).filter(AbonoGasto.id == abono_id, AbonoGasto.gasto_id == expense_id).first()
    if not abono:
        raise HTTPException(status_code=404, detail="Abono no encontrado")
    db.delete(abono)
    db.commit()
    return {"ok": True}


@router.delete("/{expense_id}")
def delete_expense(
    expense_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: Usuario = Depends(require_internal_permission("gastos.eliminar")),
):
    db_expense = db.query(Gasto).filter(Gasto.id == expense_id).first()
    if not db_expense:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    if db_expense.insumo_id is not None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se puede eliminar: el gasto tiene un insumo relacionado."
        )

    antes = snapshot(db_expense)
    db.delete(db_expense)
    db.commit()

    auditar(db, user, "gasto", expense_id, "delete", antes=antes, request=request)

    return {"message": "Gasto eliminado"}


def _validate_insumo(data: dict, db: Session):
    """Valida que el insumo exista y que la cantidad sea positiva."""
    insumo_id = data.get("insumo_id")
    cantidad = data.get("cantidad_insumo")
    if insumo_id and not db.query(Insumo).filter(Insumo.id == insumo_id).first():
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    if cantidad is not None and float(cantidad) <= 0:
        raise HTTPException(status_code=400, detail="La cantidad del insumo debe ser mayor a 0")


# ============================================================================
# ADJUNTOS (facturas, comprobantes)
# ============================================================================

def _adjunto_dict(a: GastoAdjunto) -> dict:
    return {
        "id": a.id,
        "gasto_id": a.gasto_id,
        "nombre_archivo": a.nombre_archivo,
        "mime_type": a.mime_type,
        "tamano_bytes": a.tamano_bytes,
        "subido_por": a.subido_por,
        "subido_por_nombre": None,
        "notas": a.notas,
        "created_at": a.created_at,
        "download_url": f"/api/expenses/{a.gasto_id}/adjuntos/{a.id}/download",
    }


@router.get("/adjuntos/config", dependencies=[Depends(require_internal_permission("gastos.ver"))])
def get_adjuntos_config(db: Session = Depends(get_db)):
    """Diagnóstico: ¿dónde se están guardando los adjuntos? ¿Es escribible?"""
    param = db.query(ParametroSistema).filter(ParametroSistema.clave == PARAM_DIR_KEY).first()
    configured = (param.valor or "").strip() if param else ""
    effective = configured or DEFAULT_UPLOAD_ROOT
    try:
        os.makedirs(effective, exist_ok=True)
        escribible = os.access(effective, os.W_OK)
    except OSError:
        escribible = False
    # Espacio disponible
    libre_gb = None
    try:
        st = os.statvfs(effective)
        libre_gb = round((st.f_bavail * st.f_frsize) / (1024**3), 2)
    except OSError:
        pass
    return {
        "parametro_clave": PARAM_DIR_KEY,
        "valor_configurado": configured,
        "directorio_efectivo": effective,
        "usando_default": not bool(configured),
        "default": DEFAULT_UPLOAD_ROOT,
        "existe": os.path.isdir(effective),
        "escribible": escribible,
        "espacio_libre_gb": libre_gb,
    }


@router.get("/{expense_id}/adjuntos", response_model=List[GastoAdjuntoResponse], dependencies=[Depends(require_internal_permission("gastos.ver"))])
def list_adjuntos(expense_id: int, db: Session = Depends(get_db)):
    if not db.query(Gasto.id).filter(Gasto.id == expense_id).first():
        raise HTTPException(404, "Gasto no encontrado")
    adj = (db.query(GastoAdjunto)
             .filter(GastoAdjunto.gasto_id == expense_id)
             .order_by(GastoAdjunto.created_at.desc())
             .all())
    out = []
    for a in adj:
        d = _adjunto_dict(a)
        if a.subido_por:
            u = db.query(Usuario.nombre).filter(Usuario.id == a.subido_por).first()
            d["subido_por_nombre"] = u[0] if u else None
        out.append(d)
    return out


@router.post("/{expense_id}/adjuntos", response_model=GastoAdjuntoResponse, status_code=201, dependencies=[Depends(require_internal_permission("gastos.editar"))])
async def upload_adjunto(
    expense_id: int,
    file: UploadFile = File(...),
    notas: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    user: Usuario = Depends(get_current_user),
):
    """Sube un archivo adjunto a un gasto. Max 10MB. Acepta imágenes, PDF y texto plano."""
    if not db.query(Gasto.id).filter(Gasto.id == expense_id).first():
        raise HTTPException(404, "Gasto no encontrado")

    # Validar mime type
    mime = file.content_type or ""
    if not any(mime.startswith(p) for p in ALLOWED_MIME_PREFIXES):
        raise HTTPException(400, f"Tipo de archivo no permitido: {mime}. Solo imágenes, PDF o texto.")

    # Leer (con límite)
    contents = await file.read(MAX_FILE_BYTES + 1)
    if len(contents) > MAX_FILE_BYTES:
        raise HTTPException(400, f"Archivo excede el tamaño máximo (10MB)")

    # Path: <configured_root>/YYYY/MM/uuid_ext
    upload_root = _get_upload_dir(db)
    today = date.today()
    sub_dir = os.path.join(f"{today.year:04d}", f"{today.month:02d}")
    abs_dir = os.path.join(upload_root, sub_dir)
    os.makedirs(abs_dir, exist_ok=True)

    safe_ext = ""
    if "." in (file.filename or ""):
        safe_ext = "." + file.filename.rsplit(".", 1)[1].lower()[:10]
    storage_name = f"{uuid.uuid4().hex}{safe_ext}"
    abs_path = os.path.join(abs_dir, storage_name)

    with open(abs_path, "wb") as f:
        f.write(contents)

    # Guardamos ruta ABSOLUTA — independiente de cambios futuros del parámetro
    adj = GastoAdjunto(
        gasto_id=expense_id,
        nombre_archivo=file.filename or storage_name,
        ruta_archivo=abs_path,
        mime_type=mime[:100],
        tamano_bytes=len(contents),
        subido_por=user.id,
        notas=notas,
    )
    db.add(adj); db.commit(); db.refresh(adj)
    d = _adjunto_dict(adj)
    d["subido_por_nombre"] = user.nombre
    return d


@router.get("/{expense_id}/adjuntos/{adjunto_id}/download", dependencies=[Depends(require_internal_permission("gastos.ver"))])
def download_adjunto(expense_id: int, adjunto_id: int, db: Session = Depends(get_db)):
    a = (db.query(GastoAdjunto)
           .filter(GastoAdjunto.id == adjunto_id, GastoAdjunto.gasto_id == expense_id)
           .first())
    if not a:
        raise HTTPException(404, "Adjunto no encontrado")
    abs_path = _resolve_adjunto_path(a.ruta_archivo)
    if not os.path.isfile(abs_path):
        raise HTTPException(404, f"Archivo no encontrado en el storage: {abs_path}")
    return FileResponse(abs_path, media_type=a.mime_type or "application/octet-stream",
                        filename=a.nombre_archivo)


@router.delete("/{expense_id}/adjuntos/{adjunto_id}", status_code=204, dependencies=[Depends(require_internal_permission("gastos.editar"))])
def delete_adjunto(expense_id: int, adjunto_id: int, db: Session = Depends(get_db),
                   _user: Usuario = Depends(get_current_user)):
    a = (db.query(GastoAdjunto)
           .filter(GastoAdjunto.id == adjunto_id, GastoAdjunto.gasto_id == expense_id)
           .first())
    if not a:
        raise HTTPException(404, "Adjunto no encontrado")
    abs_path = _resolve_adjunto_path(a.ruta_archivo)
    try:
        if os.path.isfile(abs_path):
            os.remove(abs_path)
    except OSError as e:
        logger.warning("No se pudo borrar archivo %s: %s", abs_path, e)
    db.delete(a); db.commit()


# ============================================================================
# WIDGETS / DASHBOARD KPIs
# ============================================================================

@router.get("/widgets", dependencies=[Depends(require_internal_permission("gastos.ver"))])
def get_widgets(
    periodo: Optional[str] = Query(None, description="YYYY-MM. Default: mes actual"),
    db: Session = Depends(get_db),
):
    """Devuelve KPIs agregados del módulo gastos para el dashboard."""
    today = date.today()
    if periodo:
        try:
            y, m = map(int, periodo.split("-"))
        except ValueError:
            raise HTTPException(400, "Formato de período inválido. Usar YYYY-MM")
    else:
        y, m = today.year, today.month

    from calendar import monthrange
    last_day = monthrange(y, m)[1]
    inicio = date(y, m, 1)
    fin = date(y, m, last_day)
    is_current = (y == today.year and m == today.month)

    # Total del mes (solo activos)
    total_mes = float(db.query(sqlfunc.coalesce(sqlfunc.sum(Gasto.valor), 0))
                        .filter(Gasto.estado != 'anulado',
                                Gasto.fecha.between(inicio, fin)).scalar())
    pagado_mes = float(db.query(sqlfunc.coalesce(sqlfunc.sum(Gasto.monto_pagado), 0))
                         .filter(Gasto.estado != 'anulado',
                                 Gasto.fecha.between(inicio, fin)).scalar())
    pendiente_mes = round(total_mes - pagado_mes, 2)

    # Gasto hoy (solo si el período es el mes actual)
    total_hoy = 0.0
    if is_current:
        total_hoy = float(db.query(sqlfunc.coalesce(sqlfunc.sum(Gasto.valor), 0))
                            .filter(Gasto.estado != 'anulado',
                                    Gasto.fecha == today).scalar())

    # Semana actual (lunes a hoy) si es mes actual
    total_semana = 0.0
    if is_current:
        weekday = today.weekday()  # lunes=0
        from datetime import timedelta
        lunes = today - timedelta(days=weekday)
        total_semana = float(db.query(sqlfunc.coalesce(sqlfunc.sum(Gasto.valor), 0))
                               .filter(Gasto.estado != 'anulado',
                                       Gasto.fecha.between(lunes, today)).scalar())

    # Directos vs Indirectos vs Sin clasificar
    # Nota: usar índices posicionales — SQLAlchemy 2.x reserva Row.t como alias interno
    # del propio Row y `r.t` no devolvería la columna.
    rows_tipo = (db.query(sqlfunc.coalesce(Gasto.tipo_costo, 'sin_clasificar').label("tipo"),
                          sqlfunc.coalesce(sqlfunc.sum(Gasto.valor), 0).label("valor"))
                   .filter(Gasto.estado != 'anulado',
                           Gasto.fecha.between(inicio, fin))
                   .group_by(sqlfunc.coalesce(Gasto.tipo_costo, 'sin_clasificar')).all())
    por_tipo_costo = {str(r[0]): float(r[1]) for r in rows_tipo}

    # Top categorías
    rows_cat = (db.query(Categoria.id, Categoria.nombre, Categoria.color, Categoria.icono,
                         sqlfunc.coalesce(sqlfunc.sum(Gasto.valor), 0).label("total"))
                  .join(Gasto, Gasto.categoria_id == Categoria.id)
                  .filter(Gasto.estado != 'anulado',
                          Gasto.fecha.between(inicio, fin))
                  .group_by(Categoria.id, Categoria.nombre, Categoria.color, Categoria.icono)
                  .order_by(sqlfunc.sum(Gasto.valor).desc())
                  .limit(10).all())
    top_categorias = [
        {"id": r.id, "nombre": r.nombre, "color": r.color, "icono": r.icono,
         "total": float(r.total), "pct": round((float(r.total) / total_mes * 100), 1) if total_mes > 0 else 0}
        for r in rows_cat
    ]

    # Tendencia 6 meses (mes actual + 5 anteriores)
    tendencia = []
    for offset in range(5, -1, -1):
        mm = m - offset
        yy = y
        while mm <= 0:
            mm += 12; yy -= 1
        ld = monthrange(yy, mm)[1]
        total = float(db.query(sqlfunc.coalesce(sqlfunc.sum(Gasto.valor), 0))
                        .filter(Gasto.estado != 'anulado',
                                Gasto.fecha.between(date(yy, mm, 1), date(yy, mm, ld))).scalar())
        tendencia.append({"year": yy, "month": mm, "total": round(total, 2)})

    # Variación vs mes anterior
    variacion_pct = 0.0
    if len(tendencia) >= 2 and tendencia[-2]["total"] > 0:
        variacion_pct = round(
            (tendencia[-1]["total"] - tendencia[-2]["total"]) / tendencia[-2]["total"] * 100, 1)

    # Alertas de sobrecosto: categorías cuyo gasto del mes supera el promedio de los 5 meses previos en +30%
    alertas = []
    for tc in top_categorias[:5]:
        prev_total = float(db.query(sqlfunc.coalesce(sqlfunc.sum(Gasto.valor), 0))
                             .filter(Gasto.estado != 'anulado',
                                     Gasto.categoria_id == tc["id"],
                                     Gasto.fecha < inicio,
                                     Gasto.fecha >= date(tendencia[0]["year"], tendencia[0]["month"], 1)).scalar())
        prev_promedio = prev_total / 5.0
        if prev_promedio > 0 and tc["total"] > prev_promedio * 1.3:
            alertas.append({
                "categoria": tc["nombre"], "icono": tc["icono"],
                "actual": tc["total"], "promedio_previo": round(prev_promedio, 2),
                "exceso_pct": round((tc["total"] - prev_promedio) / prev_promedio * 100, 1),
            })

    # Registros y cantidad de gastos
    registros = int(db.query(sqlfunc.count(Gasto.id))
                      .filter(Gasto.estado != 'anulado',
                              Gasto.fecha.between(inicio, fin)).scalar())

    dias_transcurridos = today.day if is_current else last_day
    promedio_diario = round(total_mes / max(dias_transcurridos, 1), 2)

    return {
        "periodo": f"{y:04d}-{m:02d}",
        "is_current_month": is_current,
        "total_mes": round(total_mes, 2),
        "pagado_mes": round(pagado_mes, 2),
        "pendiente_mes": pendiente_mes,
        "total_hoy": round(total_hoy, 2),
        "total_semana": round(total_semana, 2),
        "promedio_diario": promedio_diario,
        "registros": registros,
        "por_tipo_costo": por_tipo_costo,
        "top_categorias": top_categorias,
        "tendencia_6_meses": tendencia,
        "variacion_pct": variacion_pct,
        "alertas_sobrecosto": alertas,
    }


# ============================================================================
# EXPORT (CSV / XLSX)
# ============================================================================

def _build_export_query(db, start_date, end_date, categoria_id, subcategoria_id,
                        proveedor_id, estado, estado_pago):
    q = (db.query(Gasto)
           .options(
               joinedload(Gasto.proveedor),
               joinedload(Gasto.insumo),
               joinedload(Gasto.categoria_rel),
               joinedload(Gasto.subcategoria_rel),
           )
           .order_by(Gasto.fecha.desc(), Gasto.id.desc()))
    if start_date and end_date:
        q = q.filter(Gasto.fecha >= start_date, Gasto.fecha <= end_date)
    if categoria_id:    q = q.filter(Gasto.categoria_id == categoria_id)
    if subcategoria_id: q = q.filter(Gasto.subcategoria_id == subcategoria_id)
    if proveedor_id:    q = q.filter(Gasto.proveedor_id == proveedor_id)
    if estado:          q = q.filter(Gasto.estado == estado)
    if estado_pago:     q = q.filter(Gasto.estado_pago == estado_pago)
    return q


EXPORT_HEADERS = [
    "ID", "Fecha", "Proveedor", "Categoría", "Subcategoría", "Valor",
    "Pagado", "Saldo", "Estado pago", "Estado", "Insumo", "Cantidad insumo",
]


def _row_for_export(e: Gasto):
    valor = float(e.valor or 0)
    pagado = float(e.monto_pagado or 0)
    return [
        e.id,
        str(e.fecha) if e.fecha else "",
        e.proveedor.nombre if e.proveedor else "",
        e.categoria_rel.nombre if e.categoria_rel else "",
        e.subcategoria_rel.nombre if e.subcategoria_rel else "",
        valor,
        pagado,
        round(valor - pagado, 2),
        e.estado_pago or "",
        e.estado or "",
        e.insumo.nombre if e.insumo else "",
        float(e.cantidad_insumo) if e.cantidad_insumo else "",
    ]


@router.get("/export", dependencies=[Depends(require_internal_permission("gastos.ver"))])
def export_expenses(
    format: str = Query("csv", regex="^(csv|xlsx)$"),
    start_date:      Optional[str] = None,
    end_date:        Optional[str] = None,
    categoria_id:    Optional[int] = None,
    subcategoria_id: Optional[int] = None,
    proveedor_id:    Optional[int] = None,
    estado:          Optional[str] = None,
    estado_pago:     Optional[str] = None,
    db: Session = Depends(get_db),
):
    """Exporta el listado filtrado de gastos. Formato CSV o XLSX."""
    rows = _build_export_query(db, start_date, end_date, categoria_id, subcategoria_id,
                               proveedor_id, estado, estado_pago).all()
    fname_base = f"gastos_{start_date or 'todo'}_{end_date or ''}".rstrip("_")

    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(EXPORT_HEADERS)
        for e in rows:
            w.writerow(_row_for_export(e))
        buf.seek(0)
        return StreamingResponse(
            iter([buf.getvalue()]),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{fname_base}.csv"'},
        )

    # xlsx
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos"
    # Header
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="374151")
        cell.alignment = Alignment(horizontal="left", vertical="center")
    # Datos
    for e in rows:
        ws.append(_row_for_export(e))
    # Anchos
    widths = [6, 12, 24, 22, 22, 14, 14, 14, 14, 12, 22, 14]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # Stream
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname_base}.xlsx"'},
    )
